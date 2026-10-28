import openpyxl
from unificador_listas.src.parser_pdf import limpiar_precio
import re

def parse_excel_deterministic(excel_path, selected_sheets=None):
    """
    Extracción determinista para Excel usando openpyxl (data_only=True).
    Si selected_sheets es None, procesa todas.
    """
    productos_extraidos = []

    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheets_to_process = selected_sheets if selected_sheets else wb.sheetnames

        for sheet_name in sheets_to_process:
            if sheet_name not in wb.sheetnames:
                continue
            sheet = wb[sheet_name]

            # Heurística simple: buscar fila de cabecera
            header_row_idx = None
            col_map = {} # ej: "codigo": 0, "descripcion": 1, "precio": 2, "marca": 3

            for i, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if not any(row): continue

                row_str = " ".join([str(c).upper() for c in row if c])
                if "CODIGO" in row_str or "CÓDIGO" in row_str or "ARTICULO" in row_str:
                    header_row_idx = i
                    # Mapear columnas
                    for j, cell in enumerate(row):
                        if not cell: continue
                        cell_upper = str(cell).upper()
                        if "CODIGO" in cell_upper or "CÓDIGO" in cell_upper or "ARTICULO" in cell_upper:
                            if "codigo" not in col_map: col_map["codigo"] = j
                        elif "DESCRIPCION" in cell_upper or "DESCRIPCIÓN" in cell_upper or "PRODUCTO" in cell_upper:
                            if "descripcion" not in col_map: col_map["descripcion"] = j
                        elif "PRECIO" in cell_upper or "COSTO" in cell_upper or "IMPORTE" in cell_upper:
                            if "precio" not in col_map: col_map["precio"] = j
                        elif "MARCA" in cell_upper:
                            if "marca" not in col_map: col_map["marca"] = j
                    break

            if not header_row_idx or "codigo" not in col_map or "precio" not in col_map:
                continue # No se detectó estructura tabular clara, saltar (o pasar a IA)

            # Procesar filas de datos
            for i, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if i <= header_row_idx: continue
                if not any(row): continue

                cod = str(row[col_map["codigo"]]).strip() if len(row) > col_map["codigo"] and row[col_map["codigo"]] else ""
                if not cod or cod == '$' or (cod.replace('.', '').replace(',', '').isdigit() and len(cod) < 4):
                    continue

                desc = str(row[col_map["descripcion"]]).strip() if "descripcion" in col_map and len(row) > col_map["descripcion"] and row[col_map["descripcion"]] else ""
                precio_str = row[col_map["precio"]] if len(row) > col_map["precio"] else 0
                marca = str(row[col_map["marca"]]).strip() if "marca" in col_map and len(row) > col_map["marca"] and row[col_map["marca"]] else "GENERICA"

                costo = limpiar_precio(precio_str)
                if costo > 0:
                    productos_extraidos.append({
                        "codigo_proveedor": cod,
                        "descripcion": desc,
                        "marca": marca,
                        "costo": costo
                    })

    except Exception as e:
        # Aquí registraríamos el error sin romper el flujo
        pass

    return productos_extraidos
